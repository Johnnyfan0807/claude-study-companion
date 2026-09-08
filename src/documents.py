"""Secure, in-memory text extraction for supported lecture-note formats."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".pptx"}


class DocumentExtractionError(ValueError):
    """Raised when a note cannot be safely converted to text."""


@dataclass(frozen=True)
class ExtractedDocument:
    name: str
    text: str

    @property
    def characters(self) -> int:
        return len(self.text)


def _normalise_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("The text file encoding is not supported.")


def _extract_pdf(data: bytes, max_pdf_pages: int) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise DocumentExtractionError("The PDF is encrypted.") from exc
        if not unlocked:
            raise DocumentExtractionError("The PDF is encrypted.")

    pages: list[str] = []
    for index, page in enumerate(reader.pages[:max_pdf_pages], start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(f"[Page {index}]\n{page_text}")
    if len(reader.pages) > max_pdf_pages:
        pages.append(f"[Notice: only the first {max_pdf_pages} PDF pages were extracted.]")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    blocks: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _extract_pptx(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text_parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                text_parts.append(text.strip())
        if text_parts:
            slides.append(f"[Slide {index}]\n" + "\n".join(text_parts))
    return "\n\n".join(slides)


def extract_document(
    name: str,
    data: bytes,
    *,
    max_bytes: int = 8 * 1024 * 1024,
    max_pdf_pages: int = 80,
) -> ExtractedDocument:
    """Extract text without writing the uploaded file to disk."""
    safe_name = Path(name).name
    extension = Path(safe_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(
            f"{safe_name}: unsupported type. Use TXT, MD, PDF, DOCX, or PPTX."
        )
    if not data:
        raise DocumentExtractionError(f"{safe_name}: the file is empty.")
    if len(data) > max_bytes:
        limit_mb = max_bytes / (1024 * 1024)
        raise DocumentExtractionError(f"{safe_name}: file exceeds the {limit_mb:g} MB limit.")

    try:
        if extension in {".txt", ".md"}:
            text = _decode_text(data)
        elif extension == ".pdf":
            text = _extract_pdf(data, max_pdf_pages)
        elif extension == ".docx":
            text = _extract_docx(data)
        else:
            text = _extract_pptx(data)
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(
            f"{safe_name}: the file could not be read or may be corrupted."
        ) from exc

    text = _normalise_text(text)
    if not text:
        detail = " Scanned PDFs need OCR before upload." if extension == ".pdf" else ""
        raise DocumentExtractionError(f"{safe_name}: no readable text was found.{detail}")
    return ExtractedDocument(name=safe_name, text=text)


def combine_documents(documents: list[ExtractedDocument]) -> str:
    return "\n\n".join(f"===== {document.name} =====\n{document.text}" for document in documents)
