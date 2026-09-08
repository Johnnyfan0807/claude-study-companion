import io

import pytest
from docx import Document

from src.documents import DocumentExtractionError, combine_documents, extract_document


def test_extract_utf8_text() -> None:
    document = extract_document("week-1.md", b"Bayes' theorem")
    assert document.name == "week-1.md"
    assert document.text == "Bayes' theorem"


def test_extract_docx_paragraphs_and_table() -> None:
    source = Document()
    source.add_paragraph("Expected value")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Mean"
    table.cell(0, 1).text = "E[X]"
    buffer = io.BytesIO()
    source.save(buffer)

    extracted = extract_document("notes.docx", buffer.getvalue())
    assert "Expected value" in extracted.text
    assert "Mean | E[X]" in extracted.text


def test_rejects_unsupported_type() -> None:
    with pytest.raises(DocumentExtractionError, match="unsupported"):
        extract_document("notes.exe", b"not really executable")


def test_rejects_oversized_file() -> None:
    with pytest.raises(DocumentExtractionError, match="exceeds"):
        extract_document("notes.txt", b"12345", max_bytes=4)


def test_combines_files_with_source_labels() -> None:
    first = extract_document("a.txt", b"alpha")
    second = extract_document("b.txt", b"beta")
    combined = combine_documents([first, second])
    assert "===== a.txt =====" in combined
    assert "===== b.txt =====" in combined
